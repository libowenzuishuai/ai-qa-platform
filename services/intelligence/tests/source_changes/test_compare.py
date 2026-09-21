from aiqa_intelligence.doc_ingestion.runner import parse_bytes
from aiqa_intelligence.source_changes import compare_bundles


def parse_md(text: str, document_id: str):
    return parse_bytes(text.encode(), document_id, "MARKDOWN")


def test_unchanged_spans_produce_no_changes():
    body = parse_md("# Title\nAmount > 500000\n", "v1")
    other = parse_md("# Title\nAmount > 500000\n", "v2")
    report = compare_bundles("requirements/prd.md", body, other)
    assert report["changes"] == []


def test_modified_text_at_same_markdown_line():
    old = parse_md("Line one\nLine two\n", "old")
    new = parse_md("Line one\nLine two changed\n", "new")
    report = compare_bundles("doc.md", old, new)
    kinds = [c["kind"] for c in report["changes"]]
    assert "modified" in kinds
    mod = next(c for c in report["changes"] if c["kind"] == "modified")
    assert mod["old"]["quotedText"] == "Line two"
    assert mod["new"]["quotedText"] == "Line two changed"
    assert mod["old"]["locator"] == mod["new"]["locator"]


def test_insertion_shifts_lines_as_removed_and_added():
    old = parse_md("A\nB\n", "old")
    new = parse_md("A\nInserted\nB\n", "new")
    report = compare_bundles("doc.md", old, new)
    kinds = {c["kind"] for c in report["changes"]}
    assert "added" in kinds
    assert "removed" in kinds or "modified" in kinds or "uncertain" in kinds


def test_duplicate_text_in_new_version_is_uncertain():
    old = parse_md("Repeat\n", "old")
    new = parse_md("\nRepeat\nRepeat\n", "new")
    report = compare_bundles("doc.md", old, new)
    uncertain = [c for c in report["changes"] if c["kind"] == "uncertain"]
    assert uncertain
    assert any("多处重复" in (c.get("reason") or "") for c in uncertain)


def test_docx_cell_text_change_is_modified():
    from io import BytesIO

    from docx import Document

    def docx_with(cell_text: str, doc_id: str):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = cell_text
        buf = BytesIO()
        doc.save(buf)
        return parse_bytes(buf.getvalue(), doc_id, "DOCX")

    old = docx_with("<=500000", "old")
    new = docx_with(">500000", "new")
    report = compare_bundles("rules.docx", old, new)
    assert len(report["changes"]) == 1
    assert report["changes"][0]["kind"] == "modified"
    assert report["changes"][0]["old"]["quotedText"] == "<=500000"
    assert report["changes"][0]["new"]["quotedText"] == ">500000"
    assert report["changes"][0]["old"]["locator"]["kind"] == "docx-cell"


def test_unique_text_relocated_to_different_line_is_uncertain():
    old = parse_md("Amount > 500000\n", "old")
    new = parse_md("\nAmount > 500000\n", "new")
    report = compare_bundles("doc.md", old, new)
    uncertain = [c for c in report["changes"] if c["kind"] == "uncertain"]
    assert uncertain
    assert any("不同来源坐标" in (c.get("reason") or "") for c in uncertain)


def test_removed_span_when_line_deleted():
    old = parse_md("Keep\nDrop me\n", "old")
    new = parse_md("Keep\n", "new")
    report = compare_bundles("doc.md", old, new)
    removed = [c for c in report["changes"] if c["kind"] == "removed"]
    assert len(removed) == 1
    assert removed[0]["old"]["quotedText"] == "Drop me"


def test_added_span_when_line_inserted():
    old = parse_md("Keep\n", "old")
    new = parse_md("Keep\nInserted\n", "new")
    report = compare_bundles("doc.md", old, new)
    added = [c for c in report["changes"] if c["kind"] == "added"]
    assert any(c["new"]["quotedText"] == "Inserted" for c in added)


def test_format_mismatch_raises():
    old = parse_md("x", "old")
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("x")
    buf = BytesIO()
    doc.save(buf)
    new = parse_bytes(buf.getvalue(), "new", "DOCX")
    try:
        compare_bundles("x", old, new)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "format mismatch" in str(exc)


def test_quality_only_change_requires_review():
    old = parse_md('金额 500000', 'old')
    new = parse_md('金额 500000', 'new')
    new['spans'][0]['extractionQuality'] = 'LOW'
    new['coverageSummary'].update(goodSpans=0, lowSpans=1)
    report = compare_bundles('doc.md', old, new)
    assert report['changes'] and report['changes'][0]['kind'] == 'uncertain'


def test_inserted_line_does_not_claim_existing_rule_was_rewritten():
    report = compare_bundles('doc.md', parse_md('A\nB', 'old'), parse_md('X\nA\nB', 'new'))
    assert not any(c['kind'] == 'modified' for c in report['changes'])


def test_failed_bundle_is_not_a_mass_removal():
    import pytest
    old = parse_md('rule', 'old')
    new = parse_md('', 'new')
    with pytest.raises(ValueError,match='incomplete or failed'):
        compare_bundles('doc.md',old,new)


def test_same_version_and_foreign_span_rejected():
    import pytest
    from aiqa_intelligence.errors import ServiceError
    old=parse_md('rule','old')
    with pytest.raises(ValueError,match='different document versions'):
        compare_bundles('doc.md',old,old)
    new=parse_md('rule','new')
    new['spans'][0]['documentVersionId']='foreign'
    with pytest.raises(ServiceError): compare_bundles('doc.md',old,new)
