from io import BytesIO
from docx import Document
from aiqa_intelligence.doc_ingestion.runner import parse_bytes


def test_nested_cell_retains_surrounding_business_paragraphs():
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "申请人不得审批自己的申请"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "金额单位为分"
    cell.add_paragraph("超过 500000 分须主管批准")
    out = BytesIO()
    doc.save(out)
    result = parse_bytes(out.getvalue(), "doc", "DOCX")
    text = "\n".join(b["text"] for b in result["blocks"])
    assert "申请人不得审批自己的申请" in text
    assert "超过 500000 分须主管批准" in text
    assert text.count("金额单位为分") == 1


def test_header_table_and_first_page_header_are_not_silently_lost():
    from docx.shared import Inches

    doc = Document()
    doc.add_paragraph("正文")
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "首页禁止自审"
    section.header.add_table(rows=1, cols=1, width=Inches(3)).cell(0, 0).text = (
        "审批上限 500000"
    )
    out = BytesIO()
    doc.save(out)
    result = parse_bytes(out.getvalue(), "doc", "DOCX")
    text = "\n".join(b["text"] for b in result["blocks"])
    assert "首页禁止自审" in text
    assert "审批上限 500000" in text
