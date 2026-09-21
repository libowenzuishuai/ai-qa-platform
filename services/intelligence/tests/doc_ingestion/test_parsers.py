import asyncio
import hashlib
import multiprocessing
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from aiqa_intelligence.app import create_app
from aiqa_intelligence.context import RequestContext
from aiqa_intelligence.contracts.generated import (
    DocumentParseInput,
    ParsedDocumentBundle,
    VisionModelRequest,
)
from aiqa_intelligence.contracts.validation import validate_bundle
from aiqa_intelligence.doc_ingestion.runner import parse_bytes, ParseRunner
from aiqa_intelligence.doc_ingestion.service import (
    DocumentParser,
    VISION_HINT,
    VISION_SCHEMA,
    page_vision_request,
)
from aiqa_intelligence.errors import ServiceError
from aiqa_intelligence.models import Gateway
from aiqa_intelligence.storage import ArtifactReader


def parsed(data, format="MARKDOWN", document_id="doc-1"):
    body = parse_bytes(data, document_id, format)
    validate_bundle(body)
    ParsedDocumentBundle.model_validate(body)
    ids = [b["id"] for b in body["blocks"]] + [s["id"] for s in body["spans"]]
    assert len(set(ids)) == len(ids)
    for span in body["spans"]:
        if span["quotedText"] is not None:
            assert any(span["quotedText"] in b["text"] for b in body["blocks"])
    return body


def image_pdf_page(label: str) -> bytes:
    image = Image.new("RGB", (300, 100), "white")
    ImageDraw.Draw(image).text((10, 40), label, fill="black")
    buffer = BytesIO()
    image.save(buffer, "PDF")
    buffer.seek(0)
    return buffer.getvalue()


def scanned_pdf(labels: list[str]) -> bytes:
    writer = PdfWriter()
    for label in labels:
        writer.add_page(PdfReader(BytesIO(image_pdf_page(label))).pages[0])
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def table_image_pdf(rows: list[list[str]]) -> bytes:
    """Raster table for PDF_SCANNED vision path (mock in tests)."""
    import importlib.util

    path = Path(__file__).resolve().parent / "fixtures" / "drawing.py"
    spec = importlib.util.spec_from_file_location("b3_drawing", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.table_image_pdf(rows)


def docx_with_nested_table() -> bytes:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=2, cols=2)
    for row in range(2):
        for col in range(2):
            inner.cell(row, col).text = f"N{row}{col}"
    doc.add_paragraph("Below nested table")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def pdf(pages):
    writer = PdfWriter()
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    for text in pages:
        page = writer.add_blank_page(300, 200)
        if text:
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 50 150 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def image_bytes(format="PNG"):
    output = BytesIO()
    Image.new("RGB", (20, 10), "white").save(output, format=format)
    return output.getvalue()


def input_for(tmp_path, data, format):
    (tmp_path / "source").write_bytes(data)
    return DocumentParseInput(
        documentVersionId="doc-1",
        format=format,
        storageKey="source",
        checksum=hashlib.sha256(data).hexdigest(),
        fileSizeBytes=len(data),
    )


@pytest.mark.parametrize(
    "folder", ["01-explicit-prd", "02-conflict-prd", "03-missing-boundary"]
)
def test_existing_prds_keep_all_visible_lines(folder):
    root = Path(__file__).resolve().parents[4]
    data = (root / "packages/contracts/fixtures" / folder / "source.md").read_bytes()
    body = parsed(data)
    assert body["parseStatus"] == "PARSED"
    source_lines = data.decode().splitlines()
    for span in body["spans"]:
        line = source_lines[span["locator"]["startLine"] - 1]
        assert span["quotedText"] in line
    assert len(body["spans"]) == sum(bool(line.strip()) for line in source_lines)


def test_no_h1_list_and_deep_heading_preserve_punctuation():
    body = parsed(
        "前言\n- 金额超过 5000 元必须审批\n### 审批后禁止修改金额\n1. 保留 ,} 和 ,] 原文\n+ 第二项\n!!!".encode()
    )
    assert len(body["spans"]) == 6
    assert body["blocks"][3]["text"] == "保留 ,} 和 ,] 原文"
    assert body["spans"][5]["quotedText"] == "!!!"
    assert [b["kind"] for b in body["blocks"]][1:5] == [
        "listItem",
        "heading",
        "listItem",
        "listItem",
    ]


def test_fences_tables_and_txt_are_not_invented_semantics():
    data = b"```python\n# code only\n- code list\n```\n| A | B |\n|---|---|\n|1|2|"
    body = parsed(data)
    assert [b["kind"] for b in body["blocks"][:2]] == ["paragraph", "paragraph"]
    assert body["coverageSummary"]["lowSpans"] == 5
    assert body["warnings"]
    literal = parsed(data, "TXT")
    assert literal["blocks"][0]["text"] == "```python"
    assert all(b["kind"] == "paragraph" for b in literal["blocks"])


def test_docx_original_paragraph_indices_and_table_coordinates():
    doc = Document()
    doc.add_paragraph("First")
    doc.add_paragraph("")
    doc.add_paragraph("After empty")
    table = doc.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = f"R{r}C{c}"
    doc.add_paragraph("After table")
    output = BytesIO()
    doc.save(output)
    body = parsed(output.getvalue(), "DOCX")
    spans = {s["quotedText"]: s["locator"] for s in body["spans"]}
    assert spans["After empty"] == {"kind": "docx-paragraph", "paragraphIndex": 2}
    assert spans["After table"]["paragraphIndex"] == 3
    assert spans["R1C1"] == {"kind": "docx-cell", "tableIndex": 0, "row": 1, "col": 1}
    assert body["blocks"][-1]["text"] == "After table"


def test_docx_merged_cells_and_embedded_image_are_explicit():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    doc.add_picture(BytesIO(image_bytes()))
    output = BytesIO()
    doc.save(output)
    body = parsed(output.getvalue(), "DOCX")
    assert sum(b["text"] == "Merged" for b in body["blocks"]) == 1
    assert body["coverageSummary"]["unparsedSpans"] == 1
    assert any("合并" in w for w in body["warnings"])


def test_docx_horizontal_and_vertical_merged_cells_record_primary_grid():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Wide"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Tall"
    table.cell(2, 2).text = "Corner"
    output = BytesIO()
    doc.save(output)
    body = parsed(output.getvalue(), "DOCX")
    texts = {s["quotedText"]: s["locator"] for s in body["spans"] if s["quotedText"]}
    assert texts["Wide"] == {"kind": "docx-cell", "tableIndex": 0, "row": 0, "col": 0}
    assert texts["Tall"] == {"kind": "docx-cell", "tableIndex": 0, "row": 1, "col": 0}
    assert texts["Corner"] == {"kind": "docx-cell", "tableIndex": 0, "row": 2, "col": 2}
    assert sum(b["text"] == "Wide" for b in body["blocks"]) == 1
    assert sum(b["text"] == "Tall" for b in body["blocks"]) == 1
    assert any("合并" in w for w in body["warnings"])


def test_docx_footnotes_part_warns_without_parsing():
    doc = Document()
    doc.add_paragraph("Body")
    output = BytesIO()
    doc.save(output)
    data = output.getvalue()
    patched = BytesIO()
    with ZipFile(patched, "w") as out, ZipFile(BytesIO(data)) as src:
        for item in src.infolist():
            out.writestr(item, src.read(item.filename))
        out.writestr(
            "word/footnotes.xml",
            '<?xml version="1.0"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"></w:footnotes>',
        )
    body = parsed(patched.getvalue(), "DOCX")
    assert body["parseStatus"] == "PARSED"
    assert any("脚注与尾注未提取" in w for w in body["warnings"])


def test_docx_header_and_footer_use_auxiliary_paragraph_index():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "LIMIT <= 500000 FEN"
    doc.add_paragraph("Body approval rule")
    doc.sections[0].footer.paragraphs[0].text = "Page footer ,} ,]"
    output = BytesIO()
    doc.save(output)
    body = parsed(output.getvalue(), "DOCX")
    spans = {s["quotedText"]: s for s in body["spans"] if s["quotedText"]}
    assert spans["Body approval rule"]["locator"]["paragraphIndex"] == 0
    assert spans["Body approval rule"]["extractionQuality"] == "GOOD"
    header = spans["LIMIT <= 500000 FEN"]
    footer = spans["Page footer ,} ,]"]
    assert header["extractionQuality"] == "LOW"
    assert footer["extractionQuality"] == "LOW"
    assert header["locator"]["paragraphIndex"] == 1
    assert footer["locator"]["paragraphIndex"] == 2
    assert any("paragraphIndex>=1" in w for w in body["warnings"])


def test_docx_nested_table_uses_incrementing_table_index():
    body = parsed(docx_with_nested_table(), "DOCX")
    locators = {
        s["quotedText"]: s["locator"]
        for s in body["spans"]
        if s["quotedText"] and s["quotedText"].startswith("N")
    }
    assert locators["N11"] == {"kind": "docx-cell", "tableIndex": 1, "row": 1, "col": 1}
    assert all(loc["tableIndex"] == 1 for loc in locators.values())
    assert body["spans"][-1]["quotedText"] == "Below nested table"


def test_pdf_page_two_is_not_page_one():
    body = parsed(pdf(["FirstPage", "SecondPage"]), "PDF_TEXT")
    assert body["parseStatus"] == "PARSED"
    assert body["spans"][1]["quotedText"].strip() == "SecondPage"
    assert body["spans"][1]["locator"] == {"kind": "pdf-page", "page": 2}


@pytest.mark.parametrize("pages", [[None, None], ["Text", None, "ThirdPage"]])
def test_pdf_unknown_pages_remain_in_coverage(pages):
    body = parsed(pdf(pages), "PDF_TEXT")
    assert body["coverageSummary"]["unparsedSpans"] == pages.count(None)
    assert len(body["spans"]) == len(pages)
    assert body["parseStatus"] == ("PARSED" if any(pages) else "NEEDS_OCR")
    assert body["format"] == ("PDF_TEXT" if any(pages) else "PDF_SCANNED")
    assert [s["locator"]["page"] for s in body["spans"]] == list(
        range(1, len(pages) + 1)
    )


def test_pdf_page_cap_is_enforced_before_extraction():
    body = parsed(pdf([None] * 201), "PDF_TEXT")
    assert body["parseStatus"] == "FAILED"
    assert body["blocks"] == []
    assert any("200 页" in w for w in body["warnings"])


@pytest.mark.parametrize(
    "format,data",
    [
        ("PDF_TEXT", b"%PDF-garbage"),
        ("DOCX", b"not a zip"),
        ("PNG", b"\x89PNG"),
        ("MARKDOWN", b"\xff\xfe"),
        ("TXT", b"hello\0world"),
    ],
)
def test_broken_files_do_not_become_success(format, data):
    body = parsed(data, format)
    assert body["parseStatus"] == "FAILED"
    assert body["warnings"]


def test_docx_xml_entity_and_expansion_limits():
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr(
            "word/document.xml", '<!DOCTYPE foo [<!ENTITY a "attack">]><foo>&a;</foo>'
        )
    assert parsed(output.getvalue(), "DOCX")["parseStatus"] == "FAILED"
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("word/document.xml", b" " * (16 * 1024 * 1024 + 1))
    assert "16 MB" in parsed(output.getvalue(), "DOCX")["warnings"][0]


def test_long_document_id_has_stable_non_colliding_ids():
    data = b"# Title\n- one\n- two"
    body = parsed(data, document_id="a" * 128)
    assert body == parsed(data, document_id="a" * 128)
    assert max(len(s["id"]) for s in body["spans"]) < 128


class Vision:
    def __init__(self, output, error=None):
        self.output, self.error, self.calls = output, error, 0

    async def describe_image(self, request):
        return await self.describe_image_bytes(request, b"")

    async def describe_image_bytes(self, request, data):
        self.calls += 1
        if self.error:
            raise self.error
        return SimpleNamespace(parsedJson=self.output, outcome="SUCCESS")


def parse_image(tmp_path, value, *, error=None, format="PNG"):
    data = image_bytes(format)
    input = input_for(tmp_path, data, format)
    gateway = Vision(value, error)
    context = RequestContext("req", "mock", ArtifactReader(tmp_path), gateway)
    result = asyncio.run(DocumentParser().parse_document(input, context)).model_dump(
        mode="json", exclude_unset=True
    )
    validate_bundle(result)
    return result


def test_image_long_text_is_not_truncated(tmp_path):
    text = "业务原文 ,} ,] 金额 5000 元。\n" * 100
    body = parse_image(tmp_path, {"text": text})
    assert body["parseStatus"] == "PARSED"
    assert body["blocks"][0]["text"] == body["spans"][0]["quotedText"] == text
    assert body["spans"][0]["locator"] == {"kind": "image-region", "bbox": [0, 0, 1, 1]}
    assert body["coverageSummary"]["lowSpans"] == 1


@pytest.mark.parametrize(
    "value",
    [
        None,
        [],
        {"error": "failed"},
        {"text": "  "},
        {"text": 123},
        {"text": "x", "error": "bad"},
        {"text": "x" * 100001},
    ],
)
def test_invalid_vision_needs_ocr_not_business_text(tmp_path, value):
    body = parse_image(tmp_path, value, format="JPEG")
    assert body["parseStatus"] == "NEEDS_OCR"
    assert body["spans"][0]["quotedText"] is None
    assert body["coverageSummary"]["unparsedSpans"] == 1


@pytest.mark.parametrize("code", ["MODEL_NOT_CONFIGURED", "MODEL_TIMEOUT"])
def test_operational_model_errors_propagate(tmp_path, code):
    with pytest.raises(ServiceError) as error:
        parse_image(tmp_path, None, error=ServiceError(code, "not ready", 503))
    assert error.value.code == code


def test_missing_real_vision_configuration_is_503(tmp_path, monkeypatch):
    for key in ("PROVIDER", "BASE_URL", "MODEL", "API_KEY"):
        monkeypatch.delenv("AIQA_VISION_" + key, raising=False)
    input = input_for(tmp_path, image_bytes(), "PNG")
    reader = ArtifactReader(tmp_path)
    context = RequestContext("r", "real", reader, Gateway("real", reader, "test", []))
    with pytest.raises(ServiceError) as error:
        asyncio.run(DocumentParser().parse_document(input, context))
    assert error.value.code == "MODEL_NOT_CONFIGURED"


def test_http_parse_real_artifact_and_integrity_rejection(tmp_path):
    input = input_for(tmp_path, "- 金额超过 5000 元须审批".encode(), "MARKDOWN")
    request = {
        "schemaVersion": "1.0",
        "requestId": "parse-req",
        "mode": "mock",
        "timeoutMs": 10000,
        "input": input.model_dump(),
    }
    with TestClient(create_app(token="parse-test", artifact_root=tmp_path)) as client:
        assert client.get("/health").json()["capabilities"]["documentParse"] is True
        response = client.post(
            "/v1/documents/parse",
            json=request,
            headers={"authorization": "Bearer parse-test"},
        )
        assert response.status_code == 200, response.text
        assert response.json()["output"]["parseStatus"] == "PARSED"
        assert response.json()["invocations"] == []
        request["input"]["checksum"] = "0" * 64
        response = client.post(
            "/v1/documents/parse",
            json=request,
            headers={"authorization": "Bearer parse-test"},
        )
        assert response.status_code == 422
        assert response.json()["code"] == "VALIDATION_ERROR"


def test_file_size_cap(tmp_path):
    (tmp_path / "source").write_bytes(b"x" * (20 * 1024 * 1024 + 1))
    with pytest.raises(ServiceError, match="文件大小超限"):
        ArtifactReader(tmp_path).read("source")


def test_cancellation_terminates_actual_parser_child():
    before = {p.pid for p in multiprocessing.active_children()}

    async def cancel():
        runner = ParseRunner()
        task = asyncio.create_task(runner.run(pdf([None] * 200), "doc", "PDF_TEXT"))
        await asyncio.sleep(0.01)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert not ({p.pid for p in multiprocessing.active_children()} - before)
        # Slot is returned after the process exits; next job can still run.
        assert (await runner.run(b"ok", "doc", "TXT"))["parseStatus"] == "PARSED"

    asyncio.run(cancel())


def register_pdf_vision_mocks(gateway, pages: dict[int, str], input):
    for page, text in pages.items():
        gateway.register_mock(
            page_vision_request(input, page), {"text": text, "complete": True}
        )


def parse_pdf_with_ocr(tmp_path, pdf_data, pages: dict[int, str], *, format="PDF_TEXT"):
    input = input_for(tmp_path, pdf_data, format)
    gateway = Gateway("mock", ArtifactReader(tmp_path), "test", [])
    register_pdf_vision_mocks(gateway, pages, input)
    context = RequestContext("req", "mock", ArtifactReader(tmp_path), gateway)
    result = asyncio.run(DocumentParser().parse_document(input, context)).model_dump(
        mode="json", exclude_unset=True
    )
    validate_bundle(result)
    return result


def test_scanned_pdf_table_image_vision_mock_preserves_cell_text(tmp_path):
    try:
        pdf_bytes = table_image_pdf(
            [["角色", "上限"], ["申请人", "<=500000分"], ["主管", ">500000分"]]
        )
    except OSError as exc:
        pytest.skip(str(exc))
    table_text = "角色 | 上限\n申请人 | <=500000分\n主管 | >500000分"
    body = parse_pdf_with_ocr(
        tmp_path,
        pdf_bytes,
        {1: table_text},
        format="PDF_SCANNED",
    )
    assert body["parseStatus"] == "PARSED"
    assert body["format"] == "PDF_SCANNED"
    assert "<=500000" in body["blocks"][0]["text"]
    assert ">500000" in body["blocks"][0]["text"]
    assert body["spans"][0]["extractionQuality"] == "LOW"


def test_scanned_pdf_vision_preserves_source_format(tmp_path):
    body = parse_pdf_with_ocr(
        tmp_path, scanned_pdf(["金额超过 5000 元须审批"]), {1: "金额超过 5000 元须审批"}
    )
    assert body["parseStatus"] == "PARSED"
    assert body["format"] == "PDF_SCANNED"
    assert body["spans"][0]["quotedText"] == "金额超过 5000 元须审批"
    assert body["spans"][0]["extractionQuality"] == "LOW"
    assert body["spans"][0]["locator"] == {"kind": "pdf-page", "page": 1}
    assert body["coverageSummary"]["lowSpans"] == 1
    assert any("视觉大模型" in w for w in body["warnings"])


def test_mixed_pdf_vision_covers_every_page(tmp_path):
    mixed = scanned_pdf(["SecondScan"])
    writer = PdfWriter()
    writer.add_page(PdfReader(BytesIO(pdf(["FirstPage"]))).pages[0])
    writer.add_page(PdfReader(BytesIO(mixed)).pages[0])
    output = BytesIO()
    writer.write(output)
    body = parse_pdf_with_ocr(
        tmp_path, output.getvalue(), {1: "FirstPage", 2: "SecondScan"}
    )
    assert body["parseStatus"] == "PARSED"
    assert body["spans"][0]["extractionQuality"] == "LOW"
    assert body["spans"][0]["quotedText"].strip() == "FirstPage"
    assert body["spans"][1]["extractionQuality"] == "LOW"
    assert body["spans"][1]["quotedText"] == "SecondScan"
    assert body["coverageSummary"] == {
        "totalBlocks": 2,
        "goodSpans": 0,
        "lowSpans": 2,
        "unparsedSpans": 0,
    }


def test_blank_pdf_without_embedded_image_stays_needs_ocr(tmp_path):
    body = parse_pdf_with_ocr(tmp_path, pdf([None]), {})
    assert body["parseStatus"] == "NEEDS_OCR"
    assert body["coverageSummary"]["unparsedSpans"] == 1


def test_mixed_prd_sample_fixture_text_layer_without_vision():
    source = Path(__file__).resolve().parent / "fixtures" / "b3-records" / "b3-mixed-prd-sample.pdf"
    if not source.exists():
        pytest.skip("run fixtures/build_b3_mixed_fixture.py to generate sample")
    body = parsed(source.read_bytes(), "PDF_TEXT")
    assert body["parseStatus"] == "PARSED"
    assert body["format"] == "PDF_TEXT"
    pages = {s["locator"]["page"]: s for s in body["spans"] if s.get("quotedText")}
    assert "500000" in pages[1]["quotedText"]
    assert ",}" in pages[1]["quotedText"]
    assert body["coverageSummary"]["unparsedSpans"] >= 2


def test_original_b1_compressed_pdf_fixture():
    source = Path(__file__).with_name("fixtures") / "b1-two-page.pdf"
    body = parsed(source.read_bytes(), "PDF_TEXT")
    assert body["parseStatus"] == "PARSED"
    assert body["spans"][1]["quotedText"].strip() == "SecondPage"
    assert body["spans"][1]["locator"] == {"kind": "pdf-page", "page": 2}


def test_oversized_extracted_text_is_failed_not_truncated_success():
    body = parsed(b"x" * 2_000_001, "TXT")
    assert body["parseStatus"] == "FAILED"
    assert body["blocks"] == []
    assert "200 万字符" in body["warnings"][0]


def test_truncated_jpeg_and_format_mismatch():
    assert parsed(image_bytes("JPEG")[:-20], "JPEG")["parseStatus"] == "FAILED"
    assert parsed(image_bytes("PNG"), "JPEG")["parseStatus"] == "FAILED"
