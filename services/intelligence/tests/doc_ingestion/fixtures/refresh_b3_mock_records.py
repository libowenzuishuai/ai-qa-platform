"""Regenerate B3 mock eval snapshots (no paid models). Run from repo root:

  PYTHONPATH=services/intelligence/src services/intelligence/.venv/Scripts/python.exe \\
    services/intelligence/tests/doc_ingestion/fixtures/refresh_b3_mock_records.py
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

from docx import Document
from PIL import Image

from aiqa_intelligence.doc_ingestion.runner import parse_bytes

ROOT = Path(__file__).resolve().parents[5]
OUT = Path(__file__).resolve().parent / "b3-records"


def write_record(case_id: str, body: dict, *, command: str, fixture: str):
    record = {
        "caseId": case_id,
        "mode": "mock",
        "fixtureRef": fixture,
        "command": command,
        "parseStatus": body["parseStatus"],
        "format": body["format"],
        "coverageSummary": body["coverageSummary"],
        "warningCount": len(body.get("warnings") or []),
        "spanCount": len(body["spans"]),
        "blockCount": len(body["blocks"]),
    }
    path = OUT / f"{case_id}.mock.json"
    path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    print("wrote", path.relative_to(ROOT))


def docx_grid() -> bytes:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = f"R{r}C{c}"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_nested() -> bytes:
    doc = Document()
    inner = doc.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            inner.cell(r, c).text = f"N{r}{c}"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def table_image_pdf() -> bytes:
    rows = [["角色", "上限"], ["申请人", "<=500000分"]]
    width, height = 400, 20 + 40 * len(rows)
    image = Image.new("RGB", (width, height), "white")
    from PIL import ImageDraw

    draw = ImageDraw.Draw(image)
    col_width = width // 2
    for row_index, row in enumerate(rows):
        y0 = 20 + row_index * 40
        draw.line([(0, y0), (width, y0)], fill="black")
        for col_index, value in enumerate(row):
            x0 = col_index * col_width
            draw.text((x0 + 8, y0 + 12), value, fill="black")
    buf = BytesIO()
    image.save(buf, "PDF")
    return buf.getvalue()


def png_bytes() -> bytes:
    buf = BytesIO()
    Image.new("RGB", (20, 10), "white").save(buf, format="PNG")
    return buf.getvalue()


def main():
    OUT.mkdir(parents=True, exist_ok=True)

    md = (ROOT / "packages/contracts/fixtures/01-explicit-prd/source.md").read_bytes()
    write_record(
        "b3-01-markdown-prd",
        parse_bytes(md, "b3-01", "MARKDOWN"),
        command="pnpm test:doc-ingestion -k test_existing_prds",
        fixture="packages/contracts/fixtures/01-explicit-prd/source.md",
    )

    write_record(
        "b3-02-docx-grid",
        parse_bytes(docx_grid(), "b3-02", "DOCX"),
        command="pnpm test:doc-ingestion -k test_docx_original_paragraph",
        fixture="pytest-generated",
    )

    write_record(
        "b3-03-docx-merge-nested",
        parse_bytes(docx_nested(), "b3-03", "DOCX"),
        command="pnpm test:doc-ingestion -k nested_table",
        fixture="pytest-generated",
    )

    pdf = (Path(__file__).resolve().parent / "b1-two-page.pdf").read_bytes()
    write_record(
        "b3-04-pdf-text-two-page",
        parse_bytes(pdf, "b3-04", "PDF_TEXT"),
        command="pnpm test:doc-ingestion -k b1-two-page",
        fixture="tests/doc_ingestion/fixtures/b1-two-page.pdf",
    )

    write_record(
        "b3-05-pdf-scanned-table",
        parse_bytes(table_image_pdf(), "b3-05", "PDF_SCANNED"),
        command="pnpm test:doc-ingestion -k table_image",
        fixture="pytest-generated-table-image-pdf",
    )

    write_record(
        "b3-06-png-vision",
        parse_bytes(png_bytes(), "b3-06", "PNG"),
        command="pnpm test:doc-ingestion -k test_image_long_text",
        fixture="pytest-generated",
    )

    write_record(
        "b3-07-broken-input",
        parse_bytes(b"not a zip", "b3-07", "DOCX"),
        command="pnpm test:doc-ingestion -k test_broken_files",
        fixture="pytest-generated",
    )


if __name__ == "__main__":
    main()
